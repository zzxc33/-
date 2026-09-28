import sys
try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document(r'd:\代码项目\毕业设计\docs\开题报告_更新版.docx')

# Step 1: 在 Paragraph[28] (进度安排正文) 后面插入 Heading 1 标题 "七、主要参考文献"
progress_para = doc.paragraphs[28]._element
new_h7 = OxmlElement('w:p')
pPr = OxmlElement('w:pPr')
pStyle = OxmlElement('w:pStyle')
pStyle.set(qn('w:val'), 'Heading1')
pPr.append(pStyle)
new_h7.append(pPr)
run = OxmlElement('w:r')
text = OxmlElement('w:t')
text.text = '七、主要参考文献'
text.set(qn('xml:space'), 'preserve')
run.append(text)
new_h7.append(run)
progress_para.addnext(new_h7)

# Step 2: 删掉重复的 "七、主要参考文献" 标题
body = doc.element.body
para_elements = [child for child in body if child.tag == qn('w:p')]
dup_count = 0
for p in list(para_elements):
    text_content = ''.join(t.text or '' for t in p.iter(qn('w:t')))
    if text_content.strip() == '七、主要参考文献':
        dup_count += 1
        if dup_count > 1:
            p.getparent().remove(p)
print('Found H7 refs:', dup_count, 'removed:', dup_count - 1)

# Step 3: 创建表格
table = doc.add_table(rows=14, cols=4)
table.style = 'Table Grid'
headers = ['序号', '各阶段工作内容', '起讫日期', '备注']
for j, h in enumerate(headers):
    table.rows[0].cells[j].text = h

phases = [
    ['1', '毕业论文（设计）选题', '2026年9月10日——9月20日', '已完成'],
    ['2', '撰写开题报告', '2026年9月21日——9月30日', '进行中'],
    ['3', '毕业论文（设计）开题', '2026年10月1日——10月15日', '待进行'],
    ['4', '毕业论文（设计）初稿', '2026年10月16日——2027年1月31日', ''],
    ['5', '中期检查', '2027年2月1日——2月7日', ''],
    ['6', '毕业论文（设计）二稿', '2027年2月8日——3月7日', ''],
    ['7', '毕业论文（设计）定稿', '2027年3月8日——3月28日', ''],
    ['8', '毕业论文（设计）检测', '2027年3月29日——4月4日', ''],
    ['9', '毕业论文（设计）评阅', '2027年4月5日——4月11日', ''],
    ['10', '答辩资格审查', '2027年4月12日——4月16日', ''],
    ['11', '毕业论文（设计）答辩', '2027年4月17日——4月27日', ''],
    ['12', '毕业论文（设计）成绩录入', '2027年5月6日——5月8日', ''],
    ['13', '其它', '', ''],
]
for i, phase in enumerate(phases):
    for j, val in enumerate(phase):
        table.rows[i+1].cells[j].text = val

# 将表格移到 progress_para 和 new_h7 之间
tbl_elem = table._element
new_h7.addprevious(tbl_elem)
print('Table inserted')

# 保存
out = r'd:\代码项目\毕业设计\docs\开题报告_更新版.docx'
doc.save(out)
print('Saved:', out)

# 验证
print('\n=== Verification ===')
doc2 = docx.Document(out)
print('Paragraphs:', len(doc2.paragraphs), 'Tables:', len(doc2.tables))
for i in range(27, min(36, len(doc2.paragraphs))):
    p = doc2.paragraphs[i]
    print(f'  [{i}] {p.style.name:10s} | {p.text[:60]}')
if doc2.tables:
    t = doc2.tables[0]
    print(f'\nTable: {len(t.rows)} rows x {len(t.columns)} cols')
    for ri, row in enumerate(t.rows):
        cells = [c.text.strip()[:22] for c in row.cells]
        print(f'  {cells}')
