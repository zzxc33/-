"""读取用户桌面上的开题报告（正确格式版）的进度安排"""
import sys
try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    import docx

doc = docx.Document(r'c:\Users\admin\Desktop\实习证明\钟靖120230730开题报告.docx')

print('=' * 60)
print('  桌面上的开题报告 — 进度安排表')
print('=' * 60)

for ti, tbl in enumerate(doc.tables):
    print(f'\nTable {ti}: {len(tbl.rows)} rows × {len(tbl.columns)} cols')
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip() for c in row.cells]
        print(f'  {ri}: {cells}')

# 进度安排文字部分
print('\n\n进度安排正文:')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if '进度' in txt or '选题' in txt or '开题' in txt or '初稿' in txt:
        print(f'  [{i}] {txt[:100]}')
