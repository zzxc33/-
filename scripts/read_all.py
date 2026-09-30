# -*- coding: utf-8 -*-
"""读取开题报告和当前论文的参考文献/标红内容"""
import docx

# ===== 开题报告修改稿 =====
print("=" * 60)
print("📄 开题报告修改稿")
print("=" * 60)
doc = docx.Document(r'c:\Users\admin\xwechat_files\wxid_8851c1yvwxzj22_dfd4\msg\file\2026-09\开题报告_修改稿_钟靖(1).docx')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        has_red = False
        for r in p.runs:
            if r.font.color and r.font.color.rgb:
                c = str(r.font.color.rgb).upper()
                if c in ('FF0000', 'C00000', 'FF0001'):
                    has_red = True
        marker = "🔴" if has_red else "  "
        # 只打印非空的
        print(f"  [{i}] {marker} {t[:150]}")

print(f"\n总段落: {len(doc.paragraphs)}")

# ===== 当前论文定稿版 =====
print("\n" + "=" * 60)
print("📄 当前论文定稿版 - 参考文献")
print("=" * 60)
doc2 = docx.Document(r'd:\代码项目\毕业设计\钟靖120230730毕业论文_格式版.docx')
in_refs = False
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    if '参考文献' in t and p.style.name == 'Heading 1':
        in_refs = True
        continue
    if '致  谢' in t and in_refs:
        break
    if in_refs and t:
        print(f"  {t[:150]}")

# ===== 检查标红文字 =====
print("\n" + "=" * 60)
print("🔴 检查论文中的标红文字")
print("=" * 60)
red_count = 0
for i, p in enumerate(doc2.paragraphs):
    for r in p.runs:
        try:
            if r.font.color and r.font.color.rgb:
                c = str(r.font.color.rgb).upper()
                if c in ('FF0000', 'C00000', 'FF0001'):
                    if red_count < 20:
                        print(f"  段{i}: '{r.text[:60]}' (RGB={c})")
                    red_count += 1
        except:
            pass
print(f"\n标红 Run 总数: {red_count}")

# ===== 开发环境相关内容 =====
print("\n" + "=" * 60)
print("💻 当前论文中的开发环境相关内容")
print("=" * 60)
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    if any(kw in t for kw in ['开发环境', '运行环境', '硬件环境', '软件环境', '配置', 'JDK', 'Spring', 'MySQL', '数据库配置']):
        print(f"  [{i}] {t[:150]}")
