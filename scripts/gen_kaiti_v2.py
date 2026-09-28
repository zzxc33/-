# -*- coding: utf-8 -*-
"""生成开题报告 Word 文档（升级版，含双端、管理后台、小程序端）"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- 全局样式 ----------
def setup_styles():
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0.74)

    h1 = doc.styles['Heading 1']
    h1.font.name = '黑体'
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    h2 = doc.styles['Heading 2']
    h2.font.name = '黑体'
    h2.font.size = Pt(12)
    h2.font.bold = True
    h2._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

setup_styles()

def add_field(paragraph, instr_text, dirty=True):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run._r.append(fld_begin)
    run2 = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instr_text
    run2._r.append(instr)
    run3 = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    if dirty:
        fld_sep.set(qn('w:dirty'), 'true')
    run3._r.append(fld_sep)
    run4 = paragraph.add_run('（Word 打开后按 F9 更新域）')
    run5 = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run5._r.append(fld_end)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('第 ')
    run.font.size = Pt(10.5)
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    fld_run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    fld_run._r.append(fld_begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    instr_run._r.append(instr)
    sep_run = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    fld_sep.set(qn('w:dirty'), 'true')
    sep_run._r.append(fld_sep)
    num_run = paragraph.add_run('1')
    num_run.font.size = Pt(10.5)
    end_run = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    end_run._r.append(fld_end)
    tail = paragraph.add_run(' 页')
    tail.font.size = Pt(10.5)
    tail.font.name = '宋体'
    tail._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def add_toc_field(doc):
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_before = Pt(24)
    toc_title.paragraph_format.space_after = Pt(12)
    run = toc_title.add_run('目  录')
    run.font.name = '黑体'
    run.font.size = Pt(16)
    run.font.bold = True
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    toc_para = doc.add_paragraph()
    toc_para.paragraph_format.first_line_indent = Cm(0)
    add_field(toc_para, 'TOC \\o "1-3" \\h \\z \\u', dirty=True)
    doc.add_page_break()

# ---------- 封面 ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('开题报告')
run.font.size = Pt(28)
run.font.bold = True
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
title.paragraph_format.space_before = Pt(60)
title.paragraph_format.space_after = Pt(30)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('音乐资源推荐与歌单管理系统设计与实现')
run.font.size = Pt(18)
run.font.name = '黑体'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
subtitle.paragraph_format.space_after = Pt(30)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub2.add_run('—— 基于 Spring Boot 的混合推荐引擎与双端应用')
run.font.size = Pt(14)
run.font.name = '黑体'
run.font.color.rgb = None
run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
sub2.paragraph_format.space_after = Pt(60)

# ---------- 目录 ----------
add_toc_field(doc)

# ---------- 正文 ----------
def add_body(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = '宋体'
    run.font.size = Pt(12)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    return p

doc.add_heading('一、课题背景与研究意义', level=1)
add_body('近年来数字音乐产业经历了爆发式增长，全球在线音乐订阅用户已突破 5 亿，曲库规模动辄上千万首。面对如此海量的音乐资源，用户不再是主动寻找歌曲，而是被信息过载所困扰。传统的排行榜、分类浏览、关键词搜索等获取音乐的方式，要么只能反映热门趋势无法满足个性化口味，要么需要用户具备明确的搜索意图，对于"不知道自己想听什么"的场景无能为力。推荐系统正是解决这一矛盾的核心技术，它能够在没有明确查询的情况下，主动为用户推送可能感兴趣的内容。')
add_body('主流音乐平台如网易云音乐、Spotify、Apple Music 均已将推荐算法作为其核心竞争力。然而这些工业级推荐系统大多依赖大规模用户行为数据和复杂的深度学习模型，算法细节不对外公开。单一推荐算法各有其固有的缺陷：协同过滤存在冷启动问题；内容推荐缺乏多样性；深度学习模型训练成本高。如何通过混合策略扬长避短，在中小规模数据集上实现既有不错推荐效果又具备工程完整性的音乐推荐系统，是一个值得深入研究的课题。')
add_body('本课题的研究意义体现在理论和实践两个层面。在理论层面，本研究深入探讨协同过滤与内容推荐两种经典算法的实现细节，提出一种加权融合的混合推荐策略，并用新用户冷启动场景下的对比实验来验证混合策略相比单一算法的实际优势。在实践层面，本课题构建了一个覆盖用户认证、歌曲管理、歌单收藏、播放交互、评论反馈、推荐展示等完整功能链路的音乐推荐系统。与常见的单体 Web 应用不同，本系统同时提供了基于 Thymeleaf 的服务端渲染 Web 端和基于微信小程序原生框架的移动端两个前端，后端通过 REST API 同时服务于两端，实现了前后端分离与双端适配。此外，系统还内置了推荐引擎管理后台，管理员可以实时查看推荐算法的运行状态、用户分布和各算法的单独推荐结果，便于调试和效果评估。')

doc.add_heading('二、国内外研究现状', level=1)
add_body('推荐系统作为解决信息过载问题的关键技术，自 20 世纪 90 年代兴起以来已成为机器学习和数据挖掘领域的一个重要分支。协同过滤是最早被广泛应用的推荐方法，分为基于用户的 UserCF 和基于物品的 ItemCF 两种范式。Resnick 等人在 GroupLens 项目中的 UserCF 实现以及 Linden 等人在 Amazon 中应用的 ItemCF 算法奠定了协同过滤在工业界的主流地位。然而协同过滤存在明显局限，最突出的是冷启动问题和稀疏性问题。为了弥补这些不足，基于内容的推荐方法被提出，Pazzani 和 Billsus 对内容推荐系统进行了系统性的综述，指出内容推荐独立于用户行为数据、不存在新物品冷启动问题，但推荐结果容易同质化。')
add_body('为了综合两种算法的优势，混合推荐策略成为了当前的研究热点。混合方式大致可以分为加权融合、级联混合、特征组合和元学习四大类。近年来深度学习方法如 NCF、Wide & Deep、DeepFM 等被引入推荐系统，但训练成本较高，可解释性也是一大短板。在音乐推荐平台方面，Spotify 的 Discover Weekly 播放列表融合了协同过滤、音频特征分析、文本分析等多种信号，每周为超过 1 亿用户生成个性化推荐。网易云音乐则将社交关系引入推荐。这些工业级平台背后的算法细节大多不公开，但混合策略已成为主流选择。')
add_body('现有研究中存在两个明显不足。第一，学术论文大多聚焦于算法精度的对比和提升，较少关注推荐系统的完整工程实现。第二，很多毕业设计项目将"推荐系统"简化为算法原型，只实现了推荐逻辑，缺少完整的业务闭环。本课题正是瞄准这两个空白，既要深入研究三种推荐算法，也要构建功能完整、双端可用的音乐推荐应用，让推荐引擎在真实的用户交互中发挥作用。')

doc.add_heading('三、研究内容与研究目标', level=1)
add_body('本课题的研究内容涵盖推荐算法研究和系统工程实现两大维度。在推荐算法部分，系统实现了三种推荐算法模块：协同过滤（CollaborativeRecommender）、内容推荐（ContentBasedRecommender）和混合推荐（HybridRecommenderService）。协同过滤模块采用基于用户的余弦相似度计算方法，通过 UserSongInteraction 表（记录每个用户的收藏、播放、喜欢行为）构建用户-物品交互矩阵。内容推荐模块从歌曲的风格标签提取特征向量，计算候选歌曲特征与用户画像的余弦相似度。混合推荐模块采用加权融合策略，并根据用户交互数据量动态调整权重——当用户行为数据稀疏时自动提高内容推荐的权重，当用户有充足行为数据时以协同过滤为主。')
add_body('在系统工程实现方面，后端基于 Spring Boot 3.2 框架构建，集成了 Spring Security 实现用户认证和权限控制，引入 Kaptcha 验证码库。数据持久层使用 Spring Data JPA 配合 MySQL 8.0 数据库，核心实体包括用户（User）、歌曲（Song）、歌单（Playlist）、歌单歌曲关联（PlaylistSong）、评论（Comment）和用户-歌曲交互记录（UserSongInteraction）六大表。控制层分为两类：传统页面 Controller（HomeController、SongController、PlaylistController、LoginController、AdminController 等）负责 Thymeleaf 模板渲染；REST Controller（MusicApiController、AuthApiController、AdminRecommendController）负责输出 JSON 数据供小程序端和管理后台调用。')
add_body('前端采用双端架构。第一端是基于 Thymeleaf 的 Web 端，通过 Fragment 机制将侧边栏、播放器、播放列表面板等组件模块化复用，配合自定义 CSS 设计系统（深色主题）和 JavaScript 交互逻辑，实现了风格统一、交互流畅的用户界面，特别集成了支持播放/暂停、进度条、循环模式、队列管理的 Web Music Player。第二端是基于微信小程序原生框架的移动端，使用 WXML/WXSS/JS 开发，通过 wx.request 调用后端 REST API，实现了首页推荐、搜索、歌单浏览、登录注册、歌曲详情和点赞收藏等核心功能。管理后台还包含一个推荐引擎诊断页面，管理员可以在浏览器中实时查看全局统计数据、风格分布、指定用户的三种算法分别推荐结果和权重分配。')
add_body('本课题设定了以下几个研究目标：一是实现三种推荐算法并进行推荐效果对比，验证混合策略是否优于单一算法；二是在新用户冷启动场景下进行专门测试，评估内容推荐兜底策略的有效性；三是构建完整可运行的音乐推荐系统，覆盖 Web 端和微信小程序端两个前端；四是完成推荐引擎管理后台的实现，支持实时监控和算法诊断；五是进行性能优化，包括 Gzip 压缩、静态资源缓存、DOM 引用缓存等。')

doc.add_heading('四、技术路线与研究方法', level=1)
add_body('系统采用经典的分层架构，后端分为控制层（Controller）、业务层（Service）、数据访问层（Repository）和配置层（Config）四个子层。推荐引擎作为独立的业务模块（service/recommendation 包）与普通业务逻辑解耦。REST API 路径前缀统一为 /api，在 SecurityConfig 中对 /api/** 关闭 CSRF 校验并开启 CORS 跨域支持，采用简化的 Token 认证方案（内存存储，7 天有效期）。')
add_body('协同过滤算法的具体实现流程：首先将用户交互记录转换为隐式评分矩阵（收藏记 3 分、喜欢记 2 分、播放记 1 分），然后计算目标用户与每个其他用户的余弦相似度，筛选相似度最高的 K 个相似用户，汇总其交互过的歌曲去掉目标用户已交互的，按相似用户评分加权求和得到推荐得分，最终按得分降序取前 N 首。内容推荐的实现流程：先从用户历史交互歌曲中提取风格标签构建用户画像向量，再计算候选歌曲特征与用户画像的余弦相似度作为内容推荐得分。混合推荐采用加权融合，先将两种得分归一化，再按动态权重加权求和。权重根据用户交互数动态调整：交互 ≤5 条时 CF=20%、CB=50%、热门=20%、随机=10%；>5 条时 CF=50%、CB=20%、热门=15%、随机=15%。')
add_body('技术选型依据：Spring Boot 社区成熟、生态完整，与 Spring Security、Spring Data JPA、Thymeleaf 无缝集成。MySQL 天然适合多表关联查询。微信小程序原生框架比 uni-app 轻量，适合毕设项目避免编译依赖。REST API 方案让同一后端同时服务 Web 端和小程序端，实现了真正的双端适配。')

doc.add_heading('五、可行性分析', level=1)
add_body('从技术角度看，所有技术组件均有成熟的开源实现。Spring Boot、Spring Data JPA、Thymeleaf、Spring Security、Kaptcha 均为 Apache 2.0 许可。三种推荐算法均为经典方法，数学原理是线性代数中的向量相似度计算，实现难度适中。本项目已完成基础框架、三种推荐算法、用户/歌单/评论模块、Web 端、REST API、微信小程序端和推荐引擎管理后台的开发，说明技术方案完全可行。')
add_body('从数据角度看，系统采用隐式反馈数据（播放、喜欢、收藏）。初始化时通过 DataInitializer 组件加载一批模拟歌曲和用户交互数据用于演示。推荐效果评估阶段可接入公开的 Last.fm 或 MovieLens 数据集。')
add_body('从时间角度看，核心功能已全部实现，剩余工作包括系统测试、推荐算法效果评估、毕业论文撰写和答辩准备。按 16 周的毕业设计周期计算，时间安排不存在不可控风险。')

doc.add_heading('六、进度安排', level=1)
add_body('课题按照 16 周的毕业设计周期分为六个阶段。第 1 至第 3 周为文献调研与开题阶段，查阅推荐系统相关学术文献，调研主流音乐平台推荐机制，确定技术选型，完成开题报告撰写与答辩。第 4 至第 6 周为项目骨架搭建与基础功能阶段，搭建 Spring Boot 项目，设计数据库六大核心表，实现 Spring Security 认证和 Kaptcha 验证码，完成歌曲增删改查和分类浏览。第 7 至第 9 周为推荐引擎核心实现阶段，依次实现协同过滤、内容推荐和混合推荐三种算法，通过 HybridRecommenderService 对外提供统一接口。第 10 至第 11 周为业务模块完善与双端开发阶段，完成歌单系统、评论、搜索功能，使用 Thymeleaf 实现 Web 端模板和 Web Music Player 播放器，同时开发 REST API 和微信小程序端前端。第 12 至第 13 周为测试与评估阶段，进行三种推荐算法精度评估（准确率、召回率、F1 值），验证混合策略优势，完成功能测试和性能测试。第 14 至第 16 周为论文撰写与答辩准备阶段，按照学校论文格式完成撰写，制作答辩 PPT，准备系统演示。')

doc.add_heading('七、主要参考文献', level=1)
references = [
    '[1] Resnick P, Varian H R. Recommender systems[J]. Communications of the ACM, 1997, 40(3): 56-58.',
    '[2] Adomavicius G, Tuzhilin A. Toward the next generation of recommender systems: A survey of the state-of-the-art and possible extensions[J]. IEEE Transactions on Knowledge and Data Engineering, 2005, 17(6): 734-749.',
    '[3] Linden G, Smith B, York J. Amazon.com recommendations: Item-to-item collaborative filtering[J]. IEEE Internet Computing, 2003, 7(1): 76-80.',
    '[4] Pazzani M J, Billsus D. Content-based recommendation systems[M]. The Adaptive Web. Springer, 2007: 325-341.',
    '[5] Aggarwal C C. Recommender Systems: The Textbook[M]. Cham: Springer, 2016.',
    '[6] 项亮. 推荐系统实践[M]. 北京: 人民邮电出版社, 2012.',
    '[7] Ricci F, Rokach L, Shapira B. Recommender Systems Handbook[M]. 2nd ed. New York: Springer, 2015.',
    '[8] Hu Y, Koren Y, Volinsky C. Collaborative filtering for implicit feedback datasets[C]. Proceedings of the 8th IEEE International Conference on Data Mining, 2008: 263-272.',
]
for ref in references:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.74)
    run = p.add_run(ref)
    run.font.name = '宋体'
    run.font.size = Pt(10.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p.paragraph_format.line_spacing = 1.5

# ---------- 页脚页码 ----------
section = doc.sections[0]
footer = section.footer
footer.is_linked_to_previous = False
footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
add_page_number(footer_para)

# ---------- 保存 ----------
out = r'd:\代码项目\毕业设计\docs\开题报告.docx'
doc.save(out)
print(f'✅ Word 已保存: {out}')
