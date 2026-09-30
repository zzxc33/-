# -*- coding: utf-8 -*-
"""检查开发环境章节完整内容"""
import docx

doc = docx.Document(r'd:\代码项目\毕业设计\钟靖120230730毕业论文_格式版.docx')

# 找 6.1.1 开发环境章节
for i, p in enumerate(doc.paragraphs):
    if '硬件与软件环境' in p.text or '开发环境' in p.text:
        print(f"[{i}] === {p.text} ===")
        # 打印后面 15 段
        for j in range(i+1, min(i+16, len(doc.paragraphs))):
            t = doc.paragraphs[j].text.strip()
            style = doc.paragraphs[j].style.name
            if t:
                print(f"  [{j}] ({style}) {t[:150]}")
