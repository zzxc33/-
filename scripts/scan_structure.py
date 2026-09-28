"""当前开题报告完整结构扫描 — 格式重排前的基线"""
import sys
try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    import docx

from docx.oxml.ns import qn

doc = docx.Document(r'd:\代码项目\毕业设计\钟靖120230730开题报告.docx')

print('=' * 70)
print('  当前开题报告结构（重排前基线）')
print('=' * 70)

print(f'\n【段落共 {len(doc.paragraphs)} 个】')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    style = p.style.name
    # 取第一个 run 的字体
    font_name = font_size = bold = None
    if p.runs:
        r = p.runs[0]
        font_name = r.font.name
        font_size = r.font.size
        bold = r.font.bold
    # 东亚字体
    ea_font = None
    if p.runs and p.runs[0]._element.find(qn('w:rPr')) is not None:
        ea = p.runs[0]._element.find(qn('w:rPr')).find(qn('w:ea'))
        if ea is not None:
            ea_font = ea.get(qn('w:eastAsia'))
    size_str = f'{font_size.pt}pt' if font_size else 'inherit'
    print(f'  [{i:2d}] {style:12s} | font={font_name}/{ea_font} {size_str} bold={bold} | "{txt[:55]}"')

print(f'\n【表格共 {len(doc.tables)} 个】')
for ti, tbl in enumerate(doc.tables):
    print(f'\n  ── Table {ti}: {len(tbl.rows)} 行 × {len(tbl.columns)} 列 ──')
    for ri, row in enumerate(tbl.rows):
        cells_info = []
        for ci, cell in enumerate(row.cells):
            txt = cell.text.strip()[:15]
            font_name = ea_font = size = bold = None
            if cell.paragraphs and cell.paragraphs[0].runs:
                r = cell.paragraphs[0].runs[0]
                font_name = r.font.name
                size = r.font.size
                bold = r.font.bold
                rpr = r._element.find(qn('w:rPr'))
                if rpr is not None:
                    ea = rpr.find(qn('w:ea'))
                    if ea is not None:
                        ea_font = ea.get(qn('w:eastAsia'))
            size_str = f'{size.pt}pt' if size else '?'
            cells_info.append(f'[{txt}|{ea_font}/{font_name}|{size_str}|B={bold}]')
        print(f'    行{ri}: {" | ".join(cells_info)}')

print('\n✅ 扫描完成')
