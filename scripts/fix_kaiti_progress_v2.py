import sys
try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    import docx

from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document(r'd:\代码项目\毕业设计\钟靖120230730开题报告.docx')

# Step 1: 把 Paragraph[28] 恢复为原来的简洁版本（去掉上次追加的）
original_text = '本课题严格对齐南昌应用技术师范学院毕业论文（设计）进度计划，共 13 个阶段。选题工作已于 2026 年 9 月 10 日至 9 月 20 日完成，确定课题方向为「基于 Spring Boot 的混合推荐引擎智能音乐播放平台」。开题报告撰写阶段为 2026 年 9 月 21 日至 9 月 30 日，本报告正是该阶段的交付成果。开题答辩阶段安排在 2026 年 10 月 1 日至 10 月 15 日之间进行，届时将由指导教师和评阅组对课题研究方案的可行性、工作量和创新点进行审议。'
doc.paragraphs[28].text = original_text

# Step 2: 在 Paragraph[28] 后面插入两个新段落（子标题 + 技术里程碑正文）
progress_para = doc.paragraphs[28]._element

# 2a. 插入子标题段落
subheading = OxmlElement('w:p')
pPr = OxmlElement('w:pPr')
# 使用 Normal 样式但加粗 — 或者直接用 Heading4
pStyle = OxmlElement('w:pStyle')
pStyle.set(qn('w:val'), 'Heading4')
pPr.append(pStyle)
subheading.append(pPr)
run_h = OxmlElement('w:r')
text_h = OxmlElement('w:t')
text_h.text = '（一）技术实施子里程碑'
text_h.set(qn('xml:space'), 'preserve')
run_h.append(text_h)
subheading.append(run_h)
progress_para.addnext(subheading)

# 2b. 插入技术里程碑正文
milestone = OxmlElement('w:p')
run_m = OxmlElement('w:r')
text_m = OxmlElement('w:t')
text_m.text = (
    '在学校统一的进度框架下，本课题的技术实施子里程碑安排如下：开题报告答辩通过后（2026 年 10 月中旬），'
    '第一阶段（2026.10.16 — 2026.11.30，约 6 周）完成后端 Spring Boot 项目骨架搭建、数据库五大核心表设计、'
    'Spring Security 双通道鉴权实现（Web 端表单登录 + Session + CSRF、小程序端独立 JwtAuthFilter）、Kaptcha 验证码、'
    '歌曲增删改查和分类浏览；第二阶段（2026.12.1 — 2026.12.25，约 3.5 周）完成三种推荐算法模块的实现与性能优化：'
    '协同过滤（皮尔逊 + Jaccard 双相似度、UserCF + ItemCF 双模式）、内容推荐、混合推荐加权融合，消除 N+1 查询、'
    '用 Caffeine 本地缓存加速热路径；第三阶段（2026.12.26 — 2027.1.20，约 3.5 周）完成歌单管理、搜索功能、'
    '播放行为记录、推荐引擎管理后台，开发 Thymeleaf Web 端模板 + Web Music Player 播放器和微信小程序双端前端；'
    '第四阶段（2027.1.21 — 2027.1.31，约 1.5 周）完成推荐质量量化评估模块（留出法 + Precision/Recall/NDCG）'
    '和 40+ 端点自动化功能测试。上述四个子阶段与学校「初稿」阶段（2026.10.16 — 2027.1.31，共 107 天）完全对齐。'
    '初稿阶段结束后进入中期检查，随后依次完成论文二稿、定稿、检测、评阅和答辩。'
)
text_m.set(qn('xml:space'), 'preserve')
run_m.append(text_m)
milestone.append(run_m)
subheading.addnext(milestone)

# 保存
out = r'd:\代码项目\毕业设计\钟靖120230730开题报告.docx'
doc.save(out)
print('✅ 技术子里程碑已插入')

# 同步到 docs 版本
import shutil
shutil.copy2(out, r'd:\代码项目\毕业设计\docs\开题报告_更新版.docx')
print('✅ docs 版本已同步')

# 验证
print('\n=== 验证 ===')
doc2 = docx.Document(out)
for i in range(27, min(36, len(doc2.paragraphs))):
    p = doc2.paragraphs[i]
    print(f'[{i}] {p.style.name:10s} | {p.text[:70]}')
print(f'\n表格数: {len(doc2.tables)}')
if doc2.tables:
    print(f'Table 0: {len(doc2.tables[0].rows)} rows (学校 13 阶段)')
