"""修正列宽：tblGrid 的 w:w 单位是 dxa（1cm = 567 dxa）"""
import sys
try:
    import docx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    import docx

from docx.oxml.ns import qn

doc_path = r'd:\代码项目\毕业设计\钟靖120230730开题报告.docx'
doc = docx.Document(doc_path)
tbl = doc.tables[0]

# dxa 换算：1 cm ≈ 567 dxa（精确值：1 pt = 20 dxa，1 cm = 28.3465 pt × 20 = 566.93 dxa）
# 可用纸面 ≈ 16 cm（A4 21cm - 左右边距各 2.5cm）
target_cm = [1.2, 5.0, 4.0, 2.5]  # 序号/内容/日期/备注
target_dxa = [int(w * 567) for w in target_cm]
print(f'目标列宽: {target_cm} cm = {target_dxa} dxa, 总宽={sum(target_cm):.1f} cm')

tbl_grid = tbl._element.find(qn('w:tblGrid'))
cols = tbl_grid.findall(qn('w:gridCol'))
for i, col in enumerate(cols):
    col.set(qn('w:w'), str(target_dxa[i]))
    print(f'  列{i}: set to {target_cm[i]} cm ({target_dxa[i]} dxa)')

# 同步 tcPr/tcW
from docx.oxml import OxmlElement
for row in tbl.rows:
    for ci, cell in enumerate(row.cells):
        tc = cell._element
        tcPr = tc.get_or_add_tcPr()
        tcW = tcPr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW')
            tcPr.append(tcW)
        tcW.set(qn('w:w'), str(target_dxa[ci]))
        tcW.set(qn('w:type'), 'dxa')

doc.save(doc_path)
print(f'\n✅ 已保存: {doc_path}')

import shutil
shutil.copy2(doc_path, r'd:\代码项目\毕业设计\docs\开题报告_更新版.docx')
print('✅ docs 版本已同步')

# 验证
doc2 = docx.Document(doc_path)
tbl2 = doc2.tables[0]
tbl_grid2 = tbl2._element.find(qn('w:tblGrid'))
cols2 = tbl_grid2.findall(qn('w:gridCol'))
print(f'\n=== 验证 ===')
for i, col in enumerate(cols2):
    dxa_val = int(col.get(qn('w:w')))
    print(f'  列{i}: {dxa_val} dxa = {dxa_val/567:.2f} cm')
print(f'  总宽 ≈ {sum(int(c.get(qn("w:w")))/567 for c in cols2):.2f} cm')
